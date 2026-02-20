import re

from doc.config import settings
from doc.domain.bible_books import BOOK_NAMES
from doc.domain.resource_lookup import book_codes_for_lang_from_usfm_only
from docx import Document
from stet.domain.model import VerseReferenceDto, WordEntryDto
from stet.utils.util import is_valid_int


logger = settings.logger(__name__)

_RV_BLOCK_PATTERN = re.compile(r"<r>(.*?)</r>\s*<v>(.*?)</v>", re.DOTALL)
_REF_PATTERN = re.compile(r"^(.*) (\d+):([0-9,\- ]+)\s?(\(.*\))?$")


def _parse_ref_to_dto(
    reference_: str,
    lang0_code: str,
    lang1_code: str,
    lang0_book_codes_and_names: list[tuple[str, str]],
    lang1_book_codes_and_names: list[tuple[str, str]],
    book_names: dict[str, str],
    lang0_book_codes_and_names__: list[tuple[str, str]],
    source_text_with_bolding: str | None = None,
) -> VerseReferenceDto | None:
    match = _REF_PATTERN.match(reference_)
    if not match:
        logger.warning("Couldn't parse %s", reference_)
        return None
    book_name = match.group(1).replace("\n", "")
    book_codes_and_names_ = [
        (bc, bn) for bc, bn in lang0_book_codes_and_names if bn == book_name
    ]
    if not book_codes_and_names_:
        book_codes_and_names_ = [
            (bc, bn) for bc, bn in book_names.items() if bn == book_name
        ]
    book_code_and_name_ = book_codes_and_names_[0] if book_codes_and_names_ else None
    if book_code_and_name_:
        lang0_book_codes_and_names__.append(book_code_and_name_)
    chapter_num = int(match.group(2))
    verses = match.group(3)
    comment = match.group(4)
    source_reference = (
        f"{book_name} {chapter_num}:{verses}{comment}" if comment
        else f"{book_name} {chapter_num}:{verses}"
    )
    lang0_book_code = book_code_and_name_[0] if book_code_and_name_ else ""
    lang1_book_code_and_name_ = next(
        (x for x in lang1_book_codes_and_names if x[0] == lang0_book_code),
        None,
    )
    lang1_book_name = lang1_book_code_and_name_[1] if lang1_book_code_and_name_ else ""
    target_reference = f"{lang1_book_name} {chapter_num}:{verses}"
    verse_refs: list[str] = verses.split(",")
    valid_verse_refs: list[str] = []
    for verse_ref in verse_refs:
        if is_valid_int(verse_ref):
            valid_verse_refs.append(str(verse_ref))
            continue
        vm = re.match(r"(\d+)-(\d+)", verse_ref)
        if vm:
            for verse_num in range(int(vm.group(1)), int(vm.group(2)) + 1):
                valid_verse_refs.append(str(verse_num))
            continue
        logger.warning("Couldn't parse verse ref: %s", verse_ref)
    return VerseReferenceDto(
        lang0_code=lang0_code,
        lang1_code=lang1_code,
        book_code=book_code_and_name_[0] if book_code_and_name_ else "",
        book_name=book_name,
        chapter_num=chapter_num,
        source_reference=source_reference,
        target_reference=target_reference,
        verse_refs=valid_verse_refs,
        source_text_with_bolding=source_text_with_bolding,
    )


def _parse_fully_specified_column3(
    col3_text: str,
    word_entry_dto: WordEntryDto,
    lang0_code: str,
    lang1_code: str,
    lang0_book_codes_and_names: list[tuple[str, str]],
    lang1_book_codes_and_names: list[tuple[str, str]],
    book_names: dict[str, str],
    lang0_book_codes_and_names__: list[tuple[str, str]],
) -> None:
    for m in _RV_BLOCK_PATTERN.finditer(col3_text):
        ref_part = m.group(1).strip()
        verse_part = m.group(2).strip()
        dto = _parse_ref_to_dto(
            ref_part,
            lang0_code,
            lang1_code,
            lang0_book_codes_and_names,
            lang1_book_codes_and_names,
            book_names,
            lang0_book_codes_and_names__,
            source_text_with_bolding=verse_part,
        )
        if dto:
            word_entry_dto.verse_ref_dtos.append(dto)


def _parse_bible_reference_column3(
    col3_text: str,
    word_entry_dto: WordEntryDto,
    lang0_code: str,
    lang1_code: str,
    lang0_book_codes_and_names: list[tuple[str, str]],
    lang1_book_codes_and_names: list[tuple[str, str]],
    book_names: dict[str, str],
    lang0_book_codes_and_names__: list[tuple[str, str]],
) -> None:
    for reference in col3_text.split("\n"):
        reference_ = reference.strip()
        if not reference_:
            continue
        dto = _parse_ref_to_dto(
            reference_,
            lang0_code,
            lang1_code,
            lang0_book_codes_and_names,
            lang1_book_codes_and_names,
            book_names,
            lang0_book_codes_and_names__,
        )
        if dto:
            word_entry_dto.verse_ref_dtos.append(dto)


def get_word_entry_dtos(
    lang0_code: str,
    lang1_code: str,
    book_names: dict[str, str] = BOOK_NAMES,
    stet_dir: str = settings.STET_DIR,
) -> tuple[list[WordEntryDto], list[tuple[str, str]]]:
    # Build data from source doc
    word_entry_dtos: list[WordEntryDto] = []
    lang0_book_codes_and_names = book_codes_for_lang_from_usfm_only(lang0_code)
    lang1_book_codes_and_names = book_codes_for_lang_from_usfm_only(lang1_code)
    lang0_book_codes_and_names__: list[tuple[str, str]] = []
    doc = Document(f"{stet_dir}/stet_{lang0_code}.docx")
    for table in doc.tables:
        for row in table.rows:
            # Create entry item
            word_entry_dto = WordEntryDto()
            # Extract word from 1st column
            match = re.match(r"(.*)(\n)?(.*)?", row.cells[0].text)
            if not match:
                raise ValueError(f"Couldn't parse word(s): {row.cells[0].text}")
            words = match.group(1)
            word_entry_dto.words = [word.strip() for word in words.split(",")]
            raw_strongs = match.group(3)
            word_entry_dto.strongs_numbers = raw_strongs.strip()
            definition = ""
            previous_paragraph_style_name = ""
            # Get definition from 2nd column
            for paragraph in row.cells[1].paragraphs:
                text = paragraph.text.strip()
                if previous_paragraph_style_name not in (paragraph.style.name, ""):
                    definition += "\n"
                if paragraph.style.name == "List Paragraph":
                    if text:
                        definition += f"- {paragraph.text.strip()}\n"
                else:
                    if text:
                        definition += f"{paragraph.text.strip()}\n"
                previous_paragraph_style_name = paragraph.style.name
            word_entry_dto.definition = definition
            # Get verse references from 3rd column
            col3_text = row.cells[2].text.strip()
            if col3_text.startswith("<r>"):
                # Fully specified format: <r>ref</r><v>verse text with <b>bold</b></v>...
                _parse_fully_specified_column3(
                    col3_text,
                    word_entry_dto,
                    lang0_code,
                    lang1_code,
                    lang0_book_codes_and_names,
                    lang1_book_codes_and_names,
                    book_names,
                    lang0_book_codes_and_names__,
                )
            else:
                _parse_bible_reference_column3(
                    row.cells[2].text,
                    word_entry_dto,
                    lang0_code,
                    lang1_code,
                    lang0_book_codes_and_names,
                    lang1_book_codes_and_names,
                    book_names,
                    lang0_book_codes_and_names__,
                )
            # If 4th column exists, get bolded words from it
            if len(row.cells) > 3 and row.cells[3].text:
                word_entry_dto.bolded_phrases = [
                    keyword.strip() for keyword in row.cells[3].text.split(",")
                ]
            word_entry_dtos.append(word_entry_dto)
    return word_entry_dtos, list(set(lang0_book_codes_and_names__))
