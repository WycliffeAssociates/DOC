import re

from document.config import settings
from document.domain.bible_books import BOOK_NAMES
from document.stet.model import VerseReferenceDto, WordEntryDto
from document.stet.util import is_valid_int
from docx import Document  # type: ignore

logger = settings.logger(__name__)


def get_word_entry_dtos(
    lang0_code: str,
    lang1_code: str,
    book_names: dict[str, str] = BOOK_NAMES,
    stet_dir: str = settings.STET_DIR,
) -> tuple[list[WordEntryDto], list[str]]:
    # Build data from source doc
    word_entry_dtos: list[WordEntryDto] = []
    book_codes_: list[str] = []
    doc = Document(f"{stet_dir}/stet_{lang0_code}.docx")
    for table in doc.tables:
        for row in table.rows:
            # Create entry item
            word_entry_dto = WordEntryDto()
            # Extract data from word field
            match = re.match(r"(.*)(\n)?(.*)?", row.cells[0].text)
            if not match:
                raise ValueError(f"Couldn't parse word: {row.cells[0].text}")
            word = match.group(1)
            word_entry_dto.word = word
            raw_strongs = match.group(3)
            word_entry_dto.strongs_numbers = raw_strongs.strip()
            definition = ""
            previous_paragraph_style_name = ""
            for paragraph in row.cells[1].paragraphs:
                text = paragraph.text.strip()
                if previous_paragraph_style_name not in (paragraph.style.name, ""):
                    definition += "\n"
                if paragraph.style.name == "List Paragraph":
                    if text:
                        definition += f"- {paragraph.text.strip()}\n"
                else:
                    if text:
                        definition += f"{paragraph.text.strip()}\n"
                previous_paragraph_style_name = paragraph.style.name
            word_entry_dto.definition = definition
            # process verse list
            for reference in row.cells[2].text.split("\n"):
                reference_ = reference.strip()
                match = re.match(r"^(.*) (\d+):([0-9,\- ]+)\s?(\(.*\))?$", reference_)
                if not match:
                    logger.warning("Couldn't parse %s", reference_)
                    continue
                if match:
                    # Extract references
                    book_name = match.group(1)
                    book_codes = [
                        book_code
                        for book_code, book_name_ in book_names.items()
                        if book_name_ == book_name
                    ]
                    book_code = book_codes[0] if book_codes else None
                    if book_code:
                        book_codes_.append(book_code)
                    chapter_num = int(match.group(2))
                    verses = match.group(3)
                    comment = match.group(4)
                    if comment:
                        source_reference = (
                            f"{book_name} {chapter_num}:{verses}{comment}"
                        )
                    else:
                        source_reference = f"{book_name} {chapter_num}:{verses}"
                    target_reference = f"{book_name} {chapter_num}:{verses}"
                    verse_refs: list[str] = verses.split(",")
                    valid_verse_refs: list[str] = []
                    for verse_ref in verse_refs:
                        if is_valid_int(verse_ref):
                            valid_verse_refs.append(str(verse_ref))
                            continue
                        match = re.match(r"(\d+)-(\d+)", verse_ref)
                        if match:
                            start_verse = int(match.group(1))
                            end_verse = int(match.group(2))
                            verse_num = start_verse
                            while verse_num <= end_verse:
                                valid_verse_refs.append(str(verse_num))
                                verse_num += 1
                            continue
                        logger.warning("Couldn't parse verse ref: %s", verse_ref)
                    verse_reference_dto = VerseReferenceDto(
                        lang0_code=lang0_code,
                        lang1_code=lang1_code,
                        book_code=book_code,
                        book_name=book_name,
                        chapter_num=chapter_num,
                        source_reference=source_reference,
                        target_reference=target_reference,
                        verse_refs=valid_verse_refs,
                    )
                    word_entry_dto.verse_ref_dtos.append(verse_reference_dto)
            word_entry_dtos.append(word_entry_dto)
    return word_entry_dtos, list(set(book_codes_))
